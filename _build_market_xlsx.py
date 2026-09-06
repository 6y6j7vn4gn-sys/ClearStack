from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Next Lane Map"
navy = PatternFill("solid", fgColor="1B2A4A")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=12)
title_font = Font(bold=True, color="1B2A4A", name="Calibri", size=16)
thin = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC'),
)

ws['A1'] = "ClearStack empire — next-lane map (passion x cash)"
ws['A1'].font = title_font
ws.merge_cells('A1:E1')
ws['A2'] = "Updated 2026-09-06 · Money Maker · Merge after ClearStack is on its feet"
ws['A2'].font = Font(italic=True, color="666666")

headers = ["Order", "Lane", "Passion", "Cash speed", "When / why"]
for i, h in enumerate(headers, 1):
    c = ws.cell(3, i, h)
    c.fill = navy
    c.font = header_font
    c.alignment = Alignment(wrap_text=True)

rows = [
    ("NOW", "ClearStack + job apps", "High / Medium", "Fastest", "Standing business + Oct ~$20k Plan A"),
    ("Merge #1", "Fakhouri Consulting Field Docs / closeout packs", "High", "Fast after sample", "Same OCR muscle as ClearStack — one skill, second invoice; $95+/hr floor"),
    ("Merge #2", "Christian Coptic catalog + Egypt-US Bridge", "Highest heart", "Medium", "Faith + Egypt identity; wait so catalog doesn't starve cash floor"),
    ("Later", "Alexandria apparel / AI org / Cybertruck", "High vision", "Capital-heavy", "After 3 months runway"),
]
for r, row in enumerate(rows, 4):
    for c, val in enumerate(row, 1):
        cell = ws.cell(r, c, val)
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 4:
            cell.fill = PatternFill("solid", fgColor="E8F0FE")

ws['A9'] = "On its feet = public brand · PMB · business Zelle or Stripe · >=1 paid job · tithe ledger"
ws['A9'].font = Font(italic=True, size=10, color="444444")
ws.merge_cells('A9:E9')
for col, w in enumerate([12, 48, 14, 16, 55], 1):
    ws.column_dimensions[get_column_letter(col)].width = w
for r in (4, 5, 6):
    ws.row_dimensions[r].height = 45

ws2 = wb.create_sheet("Market Numbers")
ws2['A1'] = "Market snapshot (2026) — teach + compare"
ws2['A1'].font = title_font
ws2.merge_cells('A1:D1')
h2 = ["Segment", "Typical market", "Our play", "Source note"]
for i, h in enumerate(h2, 1):
    c = ws2.cell(3, i, h)
    c.fill = navy
    c.font = header_font
mrows = [
    ("Document scanning (SoCal / IE)", "$0.05-$0.18 per page · ~$100-$200 per banker's box · small jobs $500-$2,500", "Fixed packages $595 / $1,250 / $1,950 — searchable digital books, not cents-per-page", "Turn Source / Emerald Document 2026; IE shops"),
    ("Owner's rep / construction advisory (CA)", "Often 1-5% of build cost · ~$150-$300/hr advisory", "Fakhouri Consulting >=$95/hr floor; Field Docs packs share ClearStack OCR", "Terrapin / industry 2026 fee guides"),
    ("Coptic / Christian diaspora apparel", "Live DTC brands (Kairos, Coptic Central, MeemNoon, Wrth)", "Catalog + POD after ClearStack stands; Egypt-US Bridge sessions as soft funnel", "Public brand sites 2026"),
]
for r, row in enumerate(mrows, 4):
    for c, val in enumerate(row, 1):
        cell = ws2.cell(r, c, val)
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in enumerate([28, 45, 45, 40], 1):
    ws2.column_dimensions[get_column_letter(col)].width = w
for r in range(4, 7):
    ws2.row_dimensions[r].height = 60

ws3 = wb.create_sheet("Acronym Glossary")
ws3['A1'] = "Acronyms — spell out first use"
ws3['A1'].font = title_font
h3 = ["Short", "Full words", "What it means for us"]
for i, h in enumerate(h3, 1):
    c = ws3.cell(3, i, h)
    c.fill = navy
    c.font = header_font
glossary = [
    ("DBA", "Doing Business As", "Trade name ClearStack"),
    ("EIN", "Employer Identification Number", "Free IRS business tax ID"),
    ("LLC", "Limited Liability Company", "Optional company form"),
    ("FBN", "Fictitious Business Name", "Same idea as DBA (SB County)"),
    ("PMB", "Private Mail Box", "UPS Store box — never home"),
    ("OCR", "Optical Character Recognition", "Makes scans searchable"),
    ("GBP", "Google Business Profile", "Free Maps / near-me listing"),
    ("ADF", "Automatic Document Feeder", "Scanner auto-feed tray"),
    ("COGS", "Cost of Goods Sold", "USB/shipping before profit"),
    ("CS-ID", "ClearStack Job ID", "e.g. CS-20260906-01"),
    ("CTA", "Call To Action", "Ask for a quote button"),
    ("QA", "Quality Assurance", "Final check before deliver"),
    ("IE", "Inland Empire", "Our service area"),
    ("POD", "Print On Demand", "Make after order (Alexandria)"),
    ("HELOC", "Home Equity Line of Credit", "Plan B vs 401(k)"),
    ("WCAB", "Workers' Compensation Appeals Board", "Workers' comp court"),
    ("EAMS", "Electronic Adjudication Management System", "WCAB online filing"),
    ("MSC", "Mandatory Settlement Conference", "Settlement meeting date"),
    ("SKU", "Stock Keeping Unit", "Sellable package code"),
    ("SEO", "Search Engine Optimization", "Words that help Google find us"),
]
for r, row in enumerate(glossary, 4):
    for c, val in enumerate(row, 1):
        cell = ws3.cell(r, c, val)
        cell.border = thin
for col, w in enumerate([10, 42, 36], 1):
    ws3.column_dimensions[get_column_letter(col)].width = w

path = "/workspace/clearstack/ClearStack-Market-Next-Lane.xlsx"
wb.save(path)
print("saved", path)
